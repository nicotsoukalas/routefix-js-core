from __future__ import annotations

import json
import math
import shutil
import subprocess
import sys
from pathlib import Path


# ============================================================
# CONFIGURAÇÃO
# ============================================================

# Pasta onde este arquivo está:
# C:\Users\ADM\Desktop\routefix-js-core\comparador

COMPARADOR_DIR = Path(__file__).resolve().parent

# Projeto JavaScript
JS_DIR = COMPARADOR_DIR.parent

# Projeto Python original
PYTHON_DIR = Path(
    r"C:\Users\ADM\Desktop\rota\rotacerta-web"
)

# Pastas de entrada e saída
ROMANEIO_DIR = JS_DIR / "Romaneio"

RESULTADOS_PYTHON_DIR = (
    JS_DIR / "Resultados_Python"
)

RESULTADOS_JS_DIR = (
    JS_DIR / "Resultados_JS"
)

RELATORIO_PATH = (
    JS_DIR / "Comparacao_RouteFix.xlsx"
)

# Arquivo temporário usado para chamar o JS
JS_BRIDGE = JS_DIR / "_bridge_comparacao.mjs"

# Pequena tolerância para latitude/longitude
TOLERANCIA_NUMERICA = 0.0000001


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================

def log(texto=""):
    print(texto, flush=True)


def limpar_pasta(pasta):
    """
    Apaga os resultados anteriores.

    Isso é importante para que uma execução nova
    não misture arquivos antigos com arquivos novos.
    """

    pasta.mkdir(parents=True, exist_ok=True)

    for item in pasta.iterdir():

        if item.is_file() or item.is_symlink():
            item.unlink()

        elif item.is_dir():
            shutil.rmtree(item)


def normalizar_valor(valor):
    """
    Normaliza valores vindos do Excel antes da comparação.
    """

    if valor is None:
        return None

    # Trata NaN
    if isinstance(valor, float):

        if math.isnan(valor):
            return None

    if isinstance(valor, str):

        valor = (
            valor
            .replace("\xa0", " ")
            .strip()
        )

        if valor == "":
            return None

        return valor

    return valor


def valores_iguais(valor_a, valor_b):
    """
    Verifica se dois valores são iguais.

    Para números utiliza uma pequena tolerância.
    Isso evita considerar diferenças insignificantes
    de casas decimais como erro.
    """

    a = normalizar_valor(valor_a)
    b = normalizar_valor(valor_b)

    if a is None and b is None:
        return True

    if a is None or b is None:
        return False

    # Número x número
    if (
        isinstance(a, (int, float))
        and isinstance(b, (int, float))
    ):

        return math.isclose(
            float(a),
            float(b),
            rel_tol=0,
            abs_tol=TOLERANCIA_NUMERICA
        )

    # Número x texto
    if isinstance(a, (int, float)) and isinstance(b, str):

        try:

            return math.isclose(
                float(a),
                float(b.replace(",", ".")),
                rel_tol=0,
                abs_tol=TOLERANCIA_NUMERICA
            )

        except ValueError:
            pass

    # Texto x número
    if isinstance(a, str) and isinstance(b, (int, float)):

        try:

            return math.isclose(
                float(a.replace(",", ".")),
                float(b),
                rel_tol=0,
                abs_tol=TOLERANCIA_NUMERICA
            )

        except ValueError:
            pass

    return str(a).strip() == str(b).strip()


# ============================================================
# PYTHON ORIGINAL
# ============================================================

def localizar_processador_python():

    """
    Tenta localizar o pipeline Python.

    Primeiro procura a versão nova/refatorada:

        services.processamento.pipeline.executar

    Caso não exista, tenta a versão antiga:

        rota_para_rotacerta.processar
    """

    if str(PYTHON_DIR) not in sys.path:
        sys.path.insert(
            0,
            str(PYTHON_DIR)
        )

    # --------------------------------------------------------
    # Primeiro: pipeline novo
    # --------------------------------------------------------

    try:

        from services.processamento.pipeline import executar

        return (
            "services.processamento.pipeline.executar",
            executar
        )

    except Exception:
        pass

    # --------------------------------------------------------
    # Segundo: implementação antiga
    # --------------------------------------------------------

    try:

        from rota_para_rotacerta import processar

        return (
            "rota_para_rotacerta.processar",
            processar
        )

    except Exception as erro:

        raise RuntimeError(
            "\nNão foi possível localizar o processador Python.\n\n"
            "O comparador procurou por:\n"
            "1) services.processamento.pipeline.executar\n"
            "2) rota_para_rotacerta.processar\n\n"
            f"Pasta pesquisada:\n{PYTHON_DIR}\n\n"
            f"Erro:\n{erro}"
        )


def executar_python(entrada, saida):

    nome, processador = (
        localizar_processador_python()
    )

    log(
        f"    Python → {nome}"
    )

    resultado = processador(
        str(entrada),
        str(saida)
    )

    # --------------------------------------------------------
    # Pipeline novo
    # --------------------------------------------------------

    if (
        hasattr(resultado, "pacotes")
        and hasattr(resultado, "paradas")
    ):

        return {
            "pacotes": int(
                resultado.pacotes
            ),

            "paradas": int(
                resultado.paradas
            ),

            "conflitos": int(
                getattr(
                    resultado,
                    "conflitos",
                    0
                ) or 0
            ),

            "avisos": list(
                getattr(
                    resultado,
                    "avisos",
                    []
                ) or []
            )
        }

    # --------------------------------------------------------
    # Processador antigo
    # --------------------------------------------------------

    if (
        isinstance(resultado, tuple)
        and len(resultado) >= 2
    ):

        return {
            "pacotes": int(resultado[0]),
            "paradas": int(resultado[1]),
            "conflitos": 0,
            "avisos": []
        }

    raise RuntimeError(
        "O Python retornou um formato inesperado: "
        f"{type(resultado).__name__}"
    )


# ============================================================
# JAVASCRIPT
# ============================================================

def criar_bridge_js():

    """
    Cria temporariamente um pequeno arquivo Node.js.

    Ele chama diretamente:

        src/modules/pipeline.js

    Não precisamos alterar o seu código principal.
    """

    codigo = r'''
import fs from "node:fs";
import path from "node:path";
import { pathToFileURL } from "node:url";

const jsDir = process.argv[2];
const inputPath = process.argv[3];
const outputPath = process.argv[4];

const pipelinePath = path.join(
    jsDir,
    "src",
    "modules",
    "pipeline.js"
);

try {

    const pipeline = await import(
        pathToFileURL(pipelinePath).href
    );

    if (
        typeof pipeline.executar !== "function"
    ) {

        throw new Error(
            `A função executar não foi encontrada em ${pipelinePath}`
        );
    }

    const input = fs.readFileSync(
        inputPath
    );

    const resultado = await pipeline.executar(
        input,
        {
            onProgress: (etapa, progresso) => {

                if (
                    typeof progresso === "number"
                ) {

                    process.stdout.write(
                        `\r    JS: ${etapa} ${Math.round(progresso * 100)}%        `
                    );

                } else {

                    process.stdout.write(
                        `\r    JS: ${etapa}        `
                    );
                }
            }
        }
    );

    process.stdout.write("\r");

    if (
        !resultado
        || !resultado.arquivoSaida
    ) {

        throw new Error(
            "O pipeline JS não retornou arquivoSaida."
        );
    }

    fs.writeFileSync(
        outputPath,
        Buffer.from(
            resultado.arquivoSaida
        )
    );

    console.log(
        JSON.stringify({
            ok: true,

            pacotes: Number(
                resultado.pacotes ?? 0
            ),

            paradas: Number(
                resultado.paradas ?? 0
            ),

            conflitos: Number(
                resultado.conflitos ?? 0
            ),

            avisos: Array.isArray(
                resultado.avisos
            )
                ? resultado.avisos
                : []
        })
    );

} catch (erro) {

    console.error(
        erro?.stack || erro
    );

    process.exitCode = 1;
}
'''

    JS_BRIDGE.write_text(
        codigo,
        encoding="utf-8"
    )


def executar_js(entrada, saida):

    criar_bridge_js()

    comando = [
        "node",
        str(JS_BRIDGE),
        str(JS_DIR),
        str(entrada),
        str(saida)
    ]

    processo = subprocess.run(
        comando,
        cwd=str(JS_DIR),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace"
    )

    # --------------------------------------------------------
    # Mostra saída do JS
    # --------------------------------------------------------

    if processo.stdout:

        for linha in processo.stdout.splitlines():

            linha_limpa = linha.strip()

            # A última linha é o JSON
            if (
                linha_limpa.startswith("{")
                and linha_limpa.endswith("}")
            ):
                continue

            if linha_limpa:
                log(
                    "    " + linha_limpa
                )

    # --------------------------------------------------------
    # Erro
    # --------------------------------------------------------

    if processo.returncode != 0:

        raise RuntimeError(
            "Erro ao executar JavaScript:\n"
            + (
                processo.stderr.strip()
                or processo.stdout.strip()
            )
        )

    # --------------------------------------------------------
    # Localiza JSON final
    # --------------------------------------------------------

    json_linha = None

    for linha in reversed(
        processo.stdout.splitlines()
    ):

        linha = linha.strip()

        if (
            linha.startswith("{")
            and linha.endswith("}")
        ):

            json_linha = linha
            break

    if not json_linha:

        raise RuntimeError(
            "O JavaScript terminou, "
            "mas não retornou o resumo esperado."
        )

    resultado = json.loads(
        json_linha
    )

    if not resultado.get("ok"):

        raise RuntimeError(
            "O pipeline JavaScript informou falha."
        )

    return {
        "pacotes": int(
            resultado.get(
                "pacotes",
                0
            )
        ),

        "paradas": int(
            resultado.get(
                "paradas",
                0
            )
        ),

        "conflitos": int(
            resultado.get(
                "conflitos",
                0
            )
        ),

        "avisos": list(
            resultado.get(
                "avisos",
                []
            ) or []
        )
    }


# ============================================================
# LEITURA DOS XLSX
# ============================================================

def ler_xlsx(caminho):

    from openpyxl import load_workbook

    if not caminho.exists():

        raise FileNotFoundError(
            f"Arquivo não encontrado:\n{caminho}"
        )

    workbook = load_workbook(
        caminho,
        data_only=True
    )

    try:

        worksheet = workbook.active

        valores = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if not valores:
            return [], []

        cabecalho = [
            ""
            if valor is None
            else str(valor).strip()
            for valor in valores[0]
        ]

        linhas = [
            list(linha)
            for linha in valores[1:]
        ]

        return (
            cabecalho,
            linhas
        )

    finally:

        workbook.close()


# ============================================================
# COMPARAÇÃO
# ============================================================

def comparar_xlsx(
    arquivo_python,
    arquivo_js
):

    cab_python, linhas_python = (
        ler_xlsx(arquivo_python)
    )

    cab_js, linhas_js = (
        ler_xlsx(arquivo_js)
    )

    diferencas = []

    # --------------------------------------------------------
    # COMPARA CABEÇALHOS
    # --------------------------------------------------------

    maior_quantidade_colunas = max(
        len(cab_python),
        len(cab_js)
    )

    for indice in range(
        maior_quantidade_colunas
    ):

        nome_python = (
            cab_python[indice]
            if indice < len(cab_python)
            else "<ausente>"
        )

        nome_js = (
            cab_js[indice]
            if indice < len(cab_js)
            else "<ausente>"
        )

        if nome_python != nome_js:

            diferencas.append({
                "linha": 1,
                "coluna": (
                    f"coluna {indice + 1}"
                ),
                "python": nome_python,
                "javascript": nome_js
            })

    # --------------------------------------------------------
    # COMPARA LINHAS E CÉLULAS
    # --------------------------------------------------------

    maior_quantidade_linhas = max(
        len(linhas_python),
        len(linhas_js)
    )

    for indice_linha in range(
        maior_quantidade_linhas
    ):

        numero_linha_excel = (
            indice_linha + 2
        )

        linha_python = (
            linhas_python[indice_linha]
            if indice_linha < len(linhas_python)
            else []
        )

        linha_js = (
            linhas_js[indice_linha]
            if indice_linha < len(linhas_js)
            else []
        )

        maior_quantidade_colunas_linha = max(
            len(linha_python),
            len(linha_js),
            len(cab_python),
            len(cab_js)
        )

        for indice_coluna in range(
            maior_quantidade_colunas_linha
        ):

            valor_python = (
                linha_python[indice_coluna]
                if indice_coluna < len(linha_python)
                else None
            )

            valor_js = (
                linha_js[indice_coluna]
                if indice_coluna < len(linha_js)
                else None
            )

            if valores_iguais(
                valor_python,
                valor_js
            ):
                continue

            if indice_coluna < len(cab_python):

                nome_coluna = (
                    cab_python[indice_coluna]
                )

            elif indice_coluna < len(cab_js):

                nome_coluna = (
                    cab_js[indice_coluna]
                )

            else:

                nome_coluna = (
                    f"coluna {indice_coluna + 1}"
                )

            diferencas.append({
                "linha": numero_linha_excel,
                "coluna": nome_coluna,
                "python": normalizar_valor(
                    valor_python
                ),
                "javascript": normalizar_valor(
                    valor_js
                )
            })

    return {
        "linhas_python": len(
            linhas_python
        ),

        "linhas_js": len(
            linhas_js
        ),

        "colunas_python": len(
            cab_python
        ),

        "colunas_js": len(
            cab_js
        ),

        "diferencas": diferencas
    }


# ============================================================
# CRIA RELATÓRIO EXCEL
# ============================================================

def criar_relatorio(
    resumos,
    diferencas
):

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()

    # ========================================================
    # ABA RESUMO
    # ========================================================

    worksheet = workbook.active
    worksheet.title = "Resumo"

    cabecalho = [
        "Arquivo",

        "Python - Pacotes",
        "JS - Pacotes",

        "Python - Paradas",
        "JS - Paradas",

        "Python - Suspeitos",
        "JS - Suspeitos",

        "Linhas Python",
        "Linhas JS",

        "Diferenças de células",

        "Resultado"
    ]

    worksheet.append(
        cabecalho
    )

    for celula in worksheet[1]:
        celula.font = Font(
            bold=True
        )

    for item in resumos:

        worksheet.append([
            item["arquivo"],

            item.get(
                "python_pacotes"
            ),

            item.get(
                "js_pacotes"
            ),

            item.get(
                "python_paradas"
            ),

            item.get(
                "js_paradas"
            ),

            item.get(
                "python_conflitos"
            ),

            item.get(
                "js_conflitos"
            ),

            item.get(
                "linhas_python"
            ),

            item.get(
                "linhas_js"
            ),

            item.get(
                "diferencas"
            ),

            item["resultado"]
        ])

    # ========================================================
    # ABA DIFERENÇAS
    # ========================================================

    worksheet_diferencas = (
        workbook.create_sheet(
            "Diferenças"
        )
    )

    worksheet_diferencas.append([
        "Arquivo",
        "Linha Excel",
        "Coluna",
        "Valor Python",
        "Valor JavaScript"
    ])

    for celula in worksheet_diferencas[1]:
        celula.font = Font(
            bold=True
        )

    for item in diferencas:

        worksheet_diferencas.append([
            item["arquivo"],
            item["linha"],
            item["coluna"],
            item["python"],
            item["javascript"]
        ])

    # ========================================================
    # ABA CONFIGURAÇÃO
    # ========================================================

    worksheet_config = (
        workbook.create_sheet(
            "Configuração"
        )
    )

    configuracoes = [
        (
            "Projeto Python",
            str(PYTHON_DIR)
        ),

        (
            "Projeto JavaScript",
            str(JS_DIR)
        ),

        (
            "Pasta Romaneio",
            str(ROMANEIO_DIR)
        ),

        (
            "Resultados Python",
            str(RESULTADOS_PYTHON_DIR)
        ),

        (
            "Resultados JS",
            str(RESULTADOS_JS_DIR)
        ),

        (
            "Tolerância numérica",
            TOLERANCIA_NUMERICA
        )
    ]

    for chave, valor in configuracoes:

        worksheet_config.append([
            chave,
            valor
        ])

    # ========================================================
    # AJUSTA LARGURA
    # ========================================================

    for folha in workbook.worksheets:

        for coluna in folha.columns:

            maior = 0

            letra = get_column_letter(
                coluna[0].column
            )

            for celula in coluna:

                valor = (
                    ""
                    if celula.value is None
                    else str(celula.value)
                )

                maior = max(
                    maior,
                    len(valor)
                )

            folha.column_dimensions[
                letra
            ].width = min(
                max(maior + 2, 12),
                60
            )

        folha.freeze_panes = "A2"

    workbook.save(
        RELATORIO_PATH
    )


# ============================================================
# PROGRAMA PRINCIPAL
# ============================================================

def main():

    log("=" * 70)
    log(
        "       COMPARADOR ROUTEFIX - PYTHON x JAVASCRIPT"
    )
    log("=" * 70)

    log(
        f"\nPython: {PYTHON_DIR}"
    )

    log(
        f"JavaScript: {JS_DIR}"
    )

    log(
        f"Romaneio: {ROMANEIO_DIR}"
    )

    # --------------------------------------------------------
    # VALIDA PASTAS
    # --------------------------------------------------------

    if not PYTHON_DIR.exists():

        raise RuntimeError(
            "Pasta do Python não encontrada:\n"
            f"{PYTHON_DIR}"
        )

    if not JS_DIR.exists():

        raise RuntimeError(
            "Pasta do JavaScript não encontrada:\n"
            f"{JS_DIR}"
        )

    if not ROMANEIO_DIR.exists():

        raise RuntimeError(
            "Pasta Romaneio não encontrada:\n"
            f"{ROMANEIO_DIR}"
        )

    # --------------------------------------------------------
    # LOCALIZA PLANILHAS
    # --------------------------------------------------------

    arquivos = sorted(
        [
            arquivo
            for arquivo
            in ROMANEIO_DIR.iterdir()
            if (
                arquivo.is_file()
                and arquivo.suffix.lower()
                in {".xlsx", ".xls"}
            )
        ],
        key=lambda arquivo:
            arquivo.name.lower()
    )

    if not arquivos:

        log(
            "\nNenhuma planilha encontrada."
        )

        return

    # --------------------------------------------------------
    # LIMPA RESULTADOS ANTERIORES
    # --------------------------------------------------------

    limpar_pasta(
        RESULTADOS_PYTHON_DIR
    )

    limpar_pasta(
        RESULTADOS_JS_DIR
    )

    resumos = []

    todas_diferencas = []

    total = len(arquivos)

    log(
        f"\nPlanilhas encontradas: {total}\n"
    )

    try:

        # ====================================================
        # PROCESSA CADA PLANILHA
        # ====================================================

        for indice, entrada in enumerate(
            arquivos,
            start=1
        ):

            nome_base = entrada.stem

            saida_python = (
                RESULTADOS_PYTHON_DIR
                / f"{nome_base}_RouteFix.xlsx"
            )

            saida_js = (
                RESULTADOS_JS_DIR
                / f"{nome_base}_RouteFix.xlsx"
            )

            log("-" * 70)

            log(
                f"[{indice}/{total}] "
                f"{entrada.name}"
            )

            item = {

                "arquivo": entrada.name,

                "python_pacotes": None,
                "js_pacotes": None,

                "python_paradas": None,
                "js_paradas": None,

                "python_conflitos": None,
                "js_conflitos": None,

                "linhas_python": None,
                "linhas_js": None,

                "diferencas": None,

                "resultado": "ERRO"
            }

            # =================================================
            # PYTHON
            # =================================================

            try:

                log(
                    "  → Executando Python..."
                )

                resultado_python = (
                    executar_python(
                        entrada,
                        saida_python
                    )
                )

                item[
                    "python_pacotes"
                ] = resultado_python[
                    "pacotes"
                ]

                item[
                    "python_paradas"
                ] = resultado_python[
                    "paradas"
                ]

                item[
                    "python_conflitos"
                ] = resultado_python[
                    "conflitos"
                ]

                log(
                    "    ✓ "
                    f"Pacotes: "
                    f"{resultado_python['pacotes']} | "
                    f"Paradas: "
                    f"{resultado_python['paradas']} | "
                    f"Suspeitos: "
                    f"{resultado_python['conflitos']}"
                )

            except Exception as erro:

                log(
                    "    ✗ ERRO Python:"
                )

                log(
                    f"    {erro}"
                )

                resumos.append(
                    item
                )

                continue

            # =================================================
            # JAVASCRIPT
            # =================================================

            try:

                log(
                    "  → Executando JavaScript..."
                )

                resultado_js = (
                    executar_js(
                        entrada,
                        saida_js
                    )
                )

                item[
                    "js_pacotes"
                ] = resultado_js[
                    "pacotes"
                ]

                item[
                    "js_paradas"
                ] = resultado_js[
                    "paradas"
                ]

                item[
                    "js_conflitos"
                ] = resultado_js[
                    "conflitos"
                ]

                log(
                    "    ✓ "
                    f"Pacotes: "
                    f"{resultado_js['pacotes']} | "
                    f"Paradas: "
                    f"{resultado_js['paradas']} | "
                    f"Suspeitos: "
                    f"{resultado_js['conflitos']}"
                )

            except Exception as erro:

                log(
                    "    ✗ ERRO JavaScript:"
                )

                log(
                    f"    {erro}"
                )

                resumos.append(
                    item
                )

                continue

            # =================================================
            # COMPARAÇÃO
            # =================================================

            try:

                log(
                    "  → Comparando os XLSX..."
                )

                comparacao = (
                    comparar_xlsx(
                        saida_python,
                        saida_js
                    )
                )

                item[
                    "linhas_python"
                ] = comparacao[
                    "linhas_python"
                ]

                item[
                    "linhas_js"
                ] = comparacao[
                    "linhas_js"
                ]

                item[
                    "diferencas"
                ] = len(
                    comparacao[
                        "diferencas"
                    ]
                )

                # Guarda cada diferença
                for diferenca in (
                    comparacao[
                        "diferencas"
                    ]
                ):

                    todas_diferencas.append({
                        "arquivo": entrada.name,
                        **diferenca
                    })

                metricas_iguais = (

                    item[
                        "python_pacotes"
                    ]
                    ==
                    item[
                        "js_pacotes"
                    ]

                    and

                    item[
                        "python_paradas"
                    ]
                    ==
                    item[
                        "js_paradas"
                    ]

                    and

                    item[
                        "python_conflitos"
                    ]
                    ==
                    item[
                        "js_conflitos"
                    ]
                )

                planilhas_iguais = (
                    item["diferencas"] == 0
                )

                # ---------------------------------------------
                # RESULTADO
                # ---------------------------------------------

                if (
                    metricas_iguais
                    and planilhas_iguais
                ):

                    item[
                        "resultado"
                    ] = "IGUAL"

                    log(
                        "    ✓ RESULTADO: IGUAL"
                    )

                elif metricas_iguais:

                    item[
                        "resultado"
                    ] = (
                        "DIFERENÇA NAS CÉLULAS"
                    )

                    log(
                        "    ⚠ "
                        f"{item['diferencas']} "
                        "diferença(s) nas células"
                    )

                else:

                    item[
                        "resultado"
                    ] = "DIFERENÇA"

                    log(
                        "    ⚠ RESULTADO: "
                        "MÉTRICAS DIFERENTES"
                    )

            except Exception as erro:

                log(
                    "    ✗ ERRO na comparação:"
                )

                log(
                    f"    {erro}"
                )

                item[
                    "resultado"
                ] = (
                    "ERRO NA COMPARAÇÃO"
                )

            resumos.append(
                item
            )

    finally:

        # Remove o bridge temporário
        if JS_BRIDGE.exists():

            try:
                JS_BRIDGE.unlink()
            except Exception:
                pass

    # ========================================================
    # GERA RELATÓRIO
    # ========================================================

    criar_relatorio(
        resumos,
        todas_diferencas
    )

    iguais = sum(
        1
        for item in resumos
        if item["resultado"] == "IGUAL"
    )

    diferentes = (
        len(resumos) - iguais
    )

    log("\n")

    log("=" * 70)

    log(
        "                       RESUMO FINAL"
    )

    log("=" * 70)

    log(
        f"Total de planilhas : {len(resumos)}"
    )

    log(
        f"Iguais             : {iguais}"
    )

    log(
        f"Com diferenças     : {diferentes}"
    )

    log(
        "\nRelatório gerado em:"
    )

    log(
        str(RELATORIO_PATH)
    )

    log("=" * 70)


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":

    try:

        main()

    except KeyboardInterrupt:

        print(
            "\n\nProcessamento cancelado."
        )

        sys.exit(1)

    except Exception as erro:

        print(
            "\nERRO FATAL:"
        )

        print(
            erro
        )

        sys.exit(1)
